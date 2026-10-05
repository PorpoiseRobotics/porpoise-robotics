"""
diagrams_explorers.py - the drawn figures for the Pathfinder Explorers course.

Explorers is the nine-to-thirteen version of the beginner course, so these
figures say less and show more than the ones in diagrams.py: no amps, no
register names, big type, and a picture of the vehicle wherever a picture of
the vehicle will do. They are built the same way - native PowerPoint shapes,
editable, sharp in print, and readable in greyscale - and they borrow the
drawing helpers from diagrams.py so the two courses look like one family.

Each function takes a Deck, adds one slide, and returns it. Every one carries
a default speaker script, written for whoever is teaching nine-year-olds.
"""

import math

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from slidelib import (AMBER, BODY_TOP, CODE_BG, CODE_FONT, CONTENT_W, GREY,
                      INK, MARGIN_L, NAVY, RULE, TEAL, WHITE, _set_text)
from diagrams import (LIGHT_AMBER, LIGHT_GREY, LIGHT_TEAL, _arrow, _box,
                      _label, _plain_line)

# The colors a real light is showing, for the drawings of the vehicle.
LED_RED = RGBColor(0xE0, 0x22, 0x22)
LED_DIM = RGBColor(0x7A, 0x14, 0x14)
LED_ORANGE = RGBColor(0xF5, 0x8A, 0x10)
LED_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WHEEL = RGBColor(0x33, 0x33, 0x33)
EARTH = RGBColor(0x2C, 0x6E, 0xC8)
MARS = RGBColor(0xC2, 0x4A, 0x22)
SAMPLE = RGBColor(0xE6, 0xD4, 0xF5)
START_GREEN = RGBColor(0xD8, 0xF0, 0xD8)

CENTER = PP_ALIGN.CENTER


def I(value):
    return Inches(value)


def _shape(slide, kind, left, top, width, height, *, fill=WHITE, edge=NAVY,
           edge_w=1.25, dashed=False):
    """A plain shape with nothing written in it."""
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if edge is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = edge
        shape.line.width = Pt(edge_w)
        if dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    # An empty paragraph still takes the theme's 18pt unless told otherwise,
    # which the overflow lint would then measure as text.
    _set_text(shape.text_frame, [""], size=8)
    return shape


def _vehicle(slide, cx, top, width, height, front=None, back=None):
    """
    The vehicle from above, front at the top. `front` and `back` are lists of
    colors, one per drawn light, painted left to right AS SEEN FROM ABOVE -
    so the right-hand end of either list is the vehicle's right side.
    """
    left = cx - Emu(int(width / 2))
    wheel_w = I(0.24)
    wheel_h = Emu(int(height * 0.3))
    for wx in (left - wheel_w + I(0.02), left + width - I(0.02)):
        for wy in (top + I(0.08), top + height - wheel_h - I(0.08)):
            _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, wx, wy, wheel_w,
                   wheel_h, fill=WHEEL, edge=WHEEL)
    _shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, height,
           fill=LIGHT_GREY, edge=NAVY, edge_w=1.5)

    for colors, y in ((front, top - I(0.2)), (back, top + height - I(0.02))):
        if not colors:
            continue
        cell = Emu(int(width / len(colors)))
        for index, color in enumerate(colors):
            _shape(slide, MSO_SHAPE.RECTANGLE, left + cell * index, y, cell,
                   I(0.22), fill=color, edge=GREY, edge_w=0.75)


# ===================================================================
# Lesson 1
# ===================================================================

def sense_think_act(deck, title="Every robot does three things", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Point at each box on a real vehicle as you say it: the controller "
        "is how it senses you today, the computer under the top plate "
        "thinks, the wheels and lights act.",
        "Ask: does a TV remote sense, think and act? (It senses your thumb, "
        "but it doesn't decide anything - you do.) Does a robot vacuum? "
        "(Yes - it bumps, decides, and turns.) Let them argue. Arguing about "
        "the edges is how the idea sticks.",
        "Say plainly that the 'think' box is where THEIR programs will live. "
        "Everything they write in this course goes in the middle box.",
    ])

    top = BODY_TOP + I(0.35)
    width, height, gap = I(3.55), I(1.1), I(0.79)
    parts = [
        ("SENSE", "the rover's ears", ["Your controller:", "buttons and sticks"],
         LIGHT_TEAL),
        ("THINK", "the rover's brain", ["The ESP32 computer", "runs YOUR program"],
         LIGHT_AMBER),
        ("ACT", "the rover's muscles", ["4 motors turn the wheels",
                                        "32 lights light up"], LIGHT_TEAL),
    ]
    for index, (word, nickname, examples, fill) in enumerate(parts):
        left = MARGIN_L + (width + gap) * index
        _box(slide, left, top, width, height, word, fill=fill, edge=NAVY,
             size=36, bold=True, color=NAVY)
        _label(slide, left, top + height + I(0.15), width, nickname, size=22,
               bold=True, color=TEAL, align=CENTER)
        _label(slide, left, top + height + I(0.7), width, examples, size=20,
               color=INK, align=CENTER)
        if index < 2:
            mid = top + Emu(int(height / 2))
            _arrow(slide, left + width + I(0.08), mid,
                   left + width + gap - I(0.08), mid, width=3)

    deck._note(slide,
               "A ROBOT is a machine that can sense, think and act by itself. "
               "A remote-control car can act - but YOU do all its thinking.",
               "info")
    return slide


def license_course(deck, title="The driving course", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Build this on the floor before the lesson with masking tape and "
        "four plastic cups. About 10 feet by 12 feet is plenty.",
        "Walk the course on foot first, slowly, with the class watching. "
        "Nine-year-olds learn a route far better by seeing a person do it "
        "than by looking at a map.",
        "The garage is the hard part. Parking slowly and precisely is a far "
        "better test of control than speed, and it is what the license "
        "rewards.",
    ])

    arena_l, arena_t = MARGIN_L + I(0.1), BODY_TOP + I(0.25)
    arena_w, arena_h = I(8.3), I(4.1)
    _shape(slide, MSO_SHAPE.RECTANGLE, arena_l, arena_t, arena_w, arena_h,
           fill=WHITE, edge=AMBER, edge_w=3)

    start_l, start_t = arena_l + I(0.25), arena_t + arena_h - I(1.25)
    _box(slide, start_l, start_t, I(1.4), I(1.0), "START", fill=START_GREEN,
         edge=NAVY, size=20, bold=True, color=NAVY)

    cup_y = arena_t + I(1.55)
    cups = [arena_l + I(2.5) + I(1.05) * n for n in range(4)]
    for x in cups:
        _shape(slide, MSO_SHAPE.OVAL, x, cup_y, I(0.42), I(0.42),
               fill=LIGHT_AMBER, edge=AMBER, edge_w=1.5)

    garage_l = arena_l + arena_w - I(1.55)
    _box(slide, garage_l, arena_t + I(1.2), I(1.35), I(1.15), "GARAGE",
         fill=WHITE, edge=TEAL, size=20, bold=True, color=TEAL, edge_w=2.5)

    # The weave, drawn as a zigzag of short arrows between the cups.
    points = [(start_l + I(1.4), start_t + I(0.3))]
    for n, x in enumerate(cups):
        y = cup_y - I(0.45) if n % 2 == 0 else cup_y + I(0.85)
        points.append((x + I(0.21), y))
    points.append((garage_l, arena_t + I(1.75)))
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        _arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.5, dashed=True)

    right = MARGIN_L + I(8.75)
    width = CONTENT_W - I(8.75)
    _label(slide, right, BODY_TOP + I(0.25), width, "How to pass", size=24,
           bold=True, color=TEAL)
    _label(slide, right, BODY_TOP + I(0.85), width,
           ["1.  Weave through the cups",
            "2.  Park in the garage - all four wheels inside",
            "3.  Back out and drive home to START",
            "",
            "Knock over a cup? Go back to START and try again."],
           size=20, color=INK)
    return slide


# ===================================================================
# Lesson 2
# ===================================================================

def setup_and_loop(deck, title="Every program has two parts", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "This is the most important idea of the whole lesson. Slow down.",
        "Act it out. For setup(), mime getting ready once - stretch, put "
        "on a backpack. For loop(), start walking on the spot and DON'T "
        "STOP while you keep talking. Keep walking until somebody laughs.",
        "Ask: if I put the hello message in setup(), how many times will it "
        "print? (Once.) And in loop()? (Forever, or until the power goes "
        "off.)",
        "Common mix-up: kids think loop() runs once and stops. The walking "
        "is what fixes it.",
    ])

    gap = I(0.6)
    col_w = Emu(int((CONTENT_W - gap) / 2))
    top = BODY_TOP + I(0.2)

    columns = [
        ("setup()", "runs ONCE, when the rover wakes up",
         ["Like getting ready in the morning:", "wake up, get dressed, pack."],
         ["void setup() {", "  startRover();", "  setLight(0, RED);", "}"],
         LIGHT_TEAL),
        ("loop()", "runs OVER and OVER, forever",
         ["Like walking: left foot, right foot,", "left foot, right foot..."],
         ["void loop() {",
          '  Serial.println("Hello! I am a Pathfinder rover.");',
          "  delay(2000);", "}"],
         LIGHT_AMBER),
    ]
    for index, (name, when, like, code, fill) in enumerate(columns):
        left = MARGIN_L + (col_w + gap) * index
        _box(slide, left, top, col_w, I(0.75), name, fill=fill, edge=NAVY,
             size=32, bold=True, color=NAVY, font=CODE_FONT)
        _label(slide, left, top + I(0.9), col_w, when, size=22, bold=True,
               color=TEAL, align=CENTER)
        _label(slide, left, top + I(1.45), col_w, like, size=20, color=INK,
               align=CENTER)
        _box(slide, left, top + I(2.4), col_w, I(1.55), code, fill=CODE_BG,
             edge=RULE, size=14, font=CODE_FONT, shape=MSO_SHAPE.RECTANGLE,
             edge_w=0.75, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)

    # A circular arrow beside loop(), because it goes round.
    swirl = _shape(slide, MSO_SHAPE.CIRCULAR_ARROW,
                   MARGIN_L + col_w + gap + col_w - I(0.95), top + I(0.05),
                   I(0.75), I(0.65), fill=AMBER, edge=None)
    swirl.rotation = 0

    deck._note(slide,
               "The rover reads your program from top to bottom, one line at "
               "a time - and it does EXACTLY what each line says.", "info")
    return slide


# ===================================================================
# Lesson 3
# ===================================================================

def light_map(deck, title="Every light has a number", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Put a vehicle on the desk the same way round as the drawing, front "
        "toward the screen, before you say a word.",
        "Counting starts at ZERO. Ask why anybody would start at zero. "
        "(Computers do. It is strange the first time and normal by Lesson 4.)",
        "Trace the loop with your finger on the real vehicle: across the "
        "front left to right, round the right-hand corner, and back across "
        "the rear right to left. It is one long string of lights.",
        "So 0 and 31 are both on the LEFT, and 15 and 16 are both on the "
        "RIGHT. Get somebody to point at 31 before you move on.",
    ])

    row_l = MARGIN_L + I(0.95)
    row_w = I(10.3)
    cell = Emu(int(row_w / 16))
    cell_h = I(0.55)
    front_y = BODY_TOP + I(0.7)
    back_y = BODY_TOP + I(3.05)

    _label(slide, row_l, BODY_TOP + I(0.12), row_w, "FRONT of the vehicle",
           size=20, bold=True, color=NAVY, align=CENTER)

    for n in range(16):
        fill = LIGHT_AMBER if n < 8 else LIGHT_TEAL
        _box(slide, row_l + cell * n, front_y, cell, cell_h, str(n), fill=fill,
             edge=NAVY, size=18, bold=True, shape=MSO_SHAPE.RECTANGLE,
             edge_w=0.75)
    for slot in range(16):
        number = 16 + slot                         # 16 at the right, 31 at the left
        fill = LIGHT_TEAL if number <= 23 else LIGHT_AMBER
        _box(slide, row_l + row_w - cell * (slot + 1), back_y, cell, cell_h,
             str(number), fill=fill, edge=NAVY, size=18, bold=True,
             shape=MSO_SHAPE.RECTANGLE, edge_w=0.75)

    _label(slide, row_l, back_y + cell_h + I(0.12), row_w,
           "BACK of the vehicle", size=20, bold=True, color=NAVY, align=CENTER)

    # The body, with the loop drawn inside it.
    body_t = front_y + cell_h + I(0.12)
    body_b = back_y - I(0.12)
    _shape(slide, MSO_SHAPE.RECTANGLE, row_l, body_t, row_w, body_b - body_t,
           fill=WHITE, edge=RULE, edge_w=1.0)
    inset = I(0.35)
    top_y = body_t + I(0.22)
    bottom_y = body_b - I(0.22)
    _arrow(slide, row_l + inset, top_y, row_l + row_w - inset, top_y,
           color=TEAL, width=2.5)
    _arrow(slide, row_l + row_w - I(0.2), top_y, row_l + row_w - I(0.2),
           bottom_y, color=TEAL, width=2.5)
    _arrow(slide, row_l + row_w - inset, bottom_y, row_l + inset, bottom_y,
           color=TEAL, width=2.5)
    _label(slide, row_l + I(0.6), body_t + I(0.5), row_w - I(1.2),
           "One long loop: across the front, then back across the rear",
           size=20, color=INK, align=CENTER)

    _label(slide, MARGIN_L - I(0.05), front_y + I(1.2), I(0.95), "LEFT",
           size=20, bold=True, color=AMBER, align=CENTER)
    _label(slide, row_l + row_w + I(0.05), front_y + I(1.2), I(0.95), "RIGHT",
           size=20, bold=True, color=TEAL, align=CENTER)

    deck._note(slide,
               "Front: 0 to 15, left to right.  Back: 16 to 31, right to "
               "left.  So the light right behind front light 5 is back "
               "light 26.", "info")
    return slide


# ===================================================================
# Lesson 4
# ===================================================================

def for_loop_parts(deck, title="A for loop, piece by piece", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Read the line out loud as English: 'for a number starting at "
        "zero, while the number is less than thirty-two, add one each "
        "time.'",
        "Point at each callout box and the piece of code it points to.",
        "The ++ is the one that looks like a typo. It just means 'add 1'. "
        "(The language C++ is named after it - C, plus one more.)",
        "Ask: what is the LAST number it uses - 31 or 32? (31. When number "
        "becomes 32, 'less than 32' is false and the loop stops.) This is "
        "the question that tells you who has really got it.",
    ])

    code = "for (int number = 0; number < 32; number++) {"
    size = 26
    char_w = 0.55 * size / 72.0                     # Consolas, inches per char
    code_w = I(len(code) * char_w + 0.1)
    code_l = MARGIN_L + Emu(int((CONTENT_W - code_w) / 2))
    code_t = BODY_TOP + I(0.35)
    _shape(slide, MSO_SHAPE.RECTANGLE, code_l - I(0.2), code_t - I(0.12),
           code_w + I(0.4), I(0.75), fill=CODE_BG, edge=RULE, edge_w=0.75)
    _label(slide, code_l, code_t, code_w, code, size=size, bold=True,
           color=NAVY, font=CODE_FONT)

    def x_at(char):
        return code_l + I(char * char_w)

    parts = [
        (12, "START", "number begins at 0"),
        (26.5, "KEEP GOING", "while number is less than 32"),
        (38, "STEP", "add 1 each time"),
    ]
    box_w = I(3.6)
    box_t = code_t + I(1.4)
    for index, (char, word, meaning) in enumerate(parts):
        box_l = MARGIN_L + I(0.2) + (box_w + I(0.42)) * index
        _box(slide, box_l, box_t, box_w, I(1.15), [word, meaning],
             fill=LIGHT_TEAL if index != 1 else LIGHT_AMBER, edge=NAVY,
             size=20, bold=False, color=NAVY)
        _arrow(slide, box_l + Emu(int(box_w / 2)), box_t,
               x_at(char), code_t + I(0.62), color=TEAL, width=2.5)

    body_t = box_t + I(1.55)
    _box(slide, code_l, body_t, I(5.2), I(0.6),
         "  setLight(number, BLUE);", fill=CODE_BG, edge=RULE, size=22,
         font=CODE_FONT, shape=MSO_SHAPE.RECTANGLE, edge_w=0.75,
         align=PP_ALIGN.LEFT)
    _label(slide, code_l + I(5.45), body_t + I(0.05), I(5.0),
           ["runs 32 times:", "number = 0, 1, 2 ... 31"], size=20, color=INK)
    return slide


# ===================================================================
# Lesson 5
# ===================================================================

def pedal_and_coast(deck, pwm_freq=20000,
                    title="How a motor goes half speed", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "The bike story: pedal hard for a moment, then coast, then pedal "
        "again. Pedal most of the time and you go fast. Pedal a little "
        "and you go slow.",
        "The motor driver does the same thing with electricity, but "
        "thousands of times a second - far too fast to see or hear. "
        "That's why the motor just seems to run smoothly at a lower speed.",
        "Ask: which row would make the wheels go fastest? Which would "
        "make them barely move? Then: what would 0% look like? (All "
        "coasting. Stopped.)",
        "Older students may want the real name: Pulse Width Modulation, "
        "or PWM. Give it to them, but don't test on it.",
    ])

    _label(slide, MARGIN_L, BODY_TOP + I(0.05), CONTENT_W,
           "The motor driver flips the power ON and OFF {:,} times every "
           "second.".format(pwm_freq), size=22, bold=True, color=TEAL,
           align=CENTER)

    # Legend
    legend_t = BODY_TOP + I(0.72)
    _shape(slide, MSO_SHAPE.RECTANGLE, MARGIN_L + I(2.4), legend_t + I(0.06),
           I(0.4), I(0.3), fill=TEAL, edge=NAVY, edge_w=0.75)
    _label(slide, MARGIN_L + I(2.9), legend_t, I(3.4), "ON = pedal",
           size=20, color=INK)
    _shape(slide, MSO_SHAPE.RECTANGLE, MARGIN_L + I(6.4), legend_t + I(0.06),
           I(0.4), I(0.3), fill=WHITE, edge=NAVY, edge_w=0.75)
    _label(slide, MARGIN_L + I(6.9), legend_t, I(3.4), "OFF = coast",
           size=20, color=INK)

    rows = [(25, "25%  slow"), (50, "50%  medium"), (100, "100%  fast")]
    label_w = I(2.25)
    track_l = MARGIN_L + label_w + I(0.1)
    track_w = CONTENT_W - label_w - I(0.1)
    periods = 8
    period = Emu(int(track_w / periods))
    row_h = I(0.6)
    for index, (percent, text) in enumerate(rows):
        y = BODY_TOP + I(1.45) + I(1.05) * index
        _label(slide, MARGIN_L, y + I(0.1), label_w, text, size=22, bold=True,
               color=NAVY)
        for p in range(periods):
            x = track_l + period * p
            on_w = Emu(int(period * percent / 100))
            _shape(slide, MSO_SHAPE.RECTANGLE, x, y, period, row_h,
                   fill=WHITE, edge=RULE, edge_w=0.75)
            _shape(slide, MSO_SHAPE.RECTANGLE, x, y, on_w, row_h, fill=TEAL,
                   edge=NAVY, edge_w=0.75)

    deck._note(slide,
               "Your speed number picks how much of each flip is ON. More ON "
               "means faster wheels.", "info")
    return slide


def tank_turns(deck, title="No steering wheel: it turns like a tank",
               speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Look under a vehicle together first: the front wheels can't turn "
        "left or right. So how does it steer?",
        "Each arrow shows which way that SIDE'S wheels are pushing, and a "
        "longer arrow means faster.",
        "Spin right is the one that surprises people: the left side goes "
        "forward and the right side goes BACKWARD, and the vehicle turns on "
        "the spot like a figure skater.",
        "Tie it straight back to the Human Tank game they just played.",
    ])

    cases = [
        ("Straight", 50, 50, "drive(50, 50)", "goes straight"),
        ("Curve right", 50, 20, "drive(50, 20)", "curves to the right"),
        ("Spin right", 50, -50, "drive(50, -50)", "spins on the spot"),
        ("Backward", -50, -50, "drive(-50, -50)", "backs up"),
    ]
    gap = I(0.21)
    panel_w = Emu(int((CONTENT_W - gap * 3) / 4))
    top = BODY_TOP + I(0.1)
    for index, (name, left_speed, right_speed, code, result) in enumerate(cases):
        left = MARGIN_L + (panel_w + gap) * index
        cx = left + Emu(int(panel_w / 2))
        _box(slide, left, top, panel_w, I(0.55), name, fill=LIGHT_TEAL,
             edge=NAVY, size=22, bold=True, color=NAVY)

        body_t = top + I(0.95)
        body_h = I(1.55)
        _vehicle(slide, cx, body_t, I(1.05), body_h)
        _label(slide, cx - I(0.5), body_t + I(0.08), I(1.0), "front",
               size=18, color=GREY, align=CENTER)

        mid = body_t + Emu(int(body_h / 2))
        for speed, x in ((left_speed, cx - I(1.0)), (right_speed, cx + I(1.0))):
            length = I(1.4 * abs(speed) / 50.0)
            half = Emu(int(length / 2))
            color = TEAL if speed > 0 else AMBER
            if speed > 0:
                _arrow(slide, x, mid + half, x, mid - half, color=color, width=5)
            else:
                _arrow(slide, x, mid - half, x, mid + half, color=color, width=5)

        _label(slide, left, body_t + body_h + I(0.2), panel_w, code, size=20,
               bold=True, color=NAVY, align=CENTER, font=CODE_FONT)
        _label(slide, left, body_t + body_h + I(0.7), panel_w, result,
               size=20, color=INK, align=CENTER)

    deck._note(slide,
               "drive(leftSpeed, rightSpeed). Make one side faster than the "
               "other and the vehicle turns toward the SLOWER side.", "info")
    return slide


# ===================================================================
# Lesson 6
# ===================================================================

def _fit_path(points, left, top, width, height):
    """Scale and move a list of (x, y) points, in any units, into a box."""
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    span = max(max(xs) - min(xs), max(ys) - min(ys), 1e-6)
    scale = min(width / (max(xs) - min(xs) or span),
                height / (max(ys) - min(ys) or span))
    out_w = (max(xs) - min(xs)) * scale
    out_h = (max(ys) - min(ys)) * scale
    off_x = left + (width - out_w) / 2
    off_y = top + (height - out_h) / 2
    return [(Emu(int(off_x + (x - min(xs)) * scale)),
             Emu(int(off_y + (y - min(ys)) * scale))) for x, y in points]


def _turtle(turn_degrees, sides=4):
    """The corners of a path: drive 1, turn right by turn_degrees, repeat."""
    x, y, heading = 0.0, 0.0, -90.0               # Start pointing up the screen
    points = [(x, y)]
    for _ in range(sides):
        x += math.cos(math.radians(heading))
        y += math.sin(math.radians(heading))
        points.append((x, y))
        heading += turn_degrees
    return points


def square_route(deck, title="Driving a square, by the clock", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Left: what we asked for. Right: what usually happens the first "
        "time, when each corner turns a little too little.",
        "Point at the gap at the end of the right-hand path. Nothing is "
        "broken - every turn was just a bit short, and the small mistakes "
        "added up.",
        "The vehicle has no eyes. It doesn't know it missed. That's the "
        "big idea of today: a robot that drives by the clock is only as "
        "good as its numbers.",
        "Ask: if the corners were turning TOO FAR, which way would the "
        "path go wrong? (It would curl inward.)",
    ])

    half = Emu(int((CONTENT_W - I(0.6)) / 2))
    box_t = BODY_TOP + I(0.75)
    box_h = I(3.1)

    for index, (heading, turn, dashed) in enumerate(
            (("What you WANT", 90, False),
             ("What usually happens FIRST", 80, True))):
        left = MARGIN_L + (half + I(0.6)) * index
        _label(slide, left, BODY_TOP + I(0.1), half, heading, size=22,
               bold=True, color=TEAL if index == 0 else AMBER, align=CENTER)
        points = _fit_path(_turtle(turn), left + I(0.9), box_t,
                           half - I(1.8), box_h)
        for (x1, y1), (x2, y2) in zip(points, points[1:]):
            _arrow(slide, x1, y1, x2, y2, color=TEAL if index == 0 else AMBER,
                   width=3.5, dashed=dashed)
        sx, sy = points[0]
        _box(slide, sx - I(0.65), sy + I(0.12), I(1.3), I(0.5), "START",
             fill=START_GREEN, edge=NAVY, size=18, bold=True, color=NAVY)
        if index == 1:
            ex, ey = points[-1]
            _label(slide, ex + I(0.25), ey - I(0.2), I(2.2), "ends up here!",
                   size=18, bold=True, color=AMBER)

    deck._note(slide,
               "The vehicle can't see where it is. Tiny mistakes add up. "
               "Change TURN_TIME by 50 at a time until it comes home.",
               "info")
    return slide


# ===================================================================
# Lesson 7
# ===================================================================

def mars_delay(deck, title="Why Mars rovers drive themselves", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Radio travels at the speed of light, which is unbelievably fast - "
        "but Mars is unbelievably far away. Depending on where the two "
        "planets are in their orbits, a message takes somewhere between "
        "about 3 and 22 minutes to get there, one way.",
        "Ask: if the rover is about to drive into a crater and you press "
        "STOP, when does the rover find out? (Minutes later.) And when do "
        "you find out it stopped? (Minutes after THAT.)",
        "Mars rovers are driven from NASA's Jet Propulsion Laboratory in "
        "Pasadena, a couple of hours' drive north of us. The team plans a "
        "whole day of moves, sends the plan, and the rover carries it out "
        "on its own. That is exactly what your vehicle is about to do.",
    ])

    mid_y = BODY_TOP + I(1.55)
    earth_d, mars_d = I(1.8), I(1.15)
    earth_l = MARGIN_L + I(0.6)
    mars_l = MARGIN_L + CONTENT_W - I(0.6) - mars_d
    _shape(slide, MSO_SHAPE.OVAL, earth_l, mid_y - Emu(int(earth_d / 2)),
           earth_d, earth_d, fill=EARTH, edge=NAVY, edge_w=1.5)
    _shape(slide, MSO_SHAPE.OVAL, mars_l, mid_y - Emu(int(mars_d / 2)),
           mars_d, mars_d, fill=MARS, edge=NAVY, edge_w=1.5)

    _label(slide, earth_l - I(0.6), mid_y + I(1.05), earth_d + I(1.2),
           ["EARTH", "the mission team"], size=20, bold=True, color=NAVY,
           align=CENTER)
    _label(slide, mars_l - I(0.6), mid_y + I(1.05), mars_d + I(1.2),
           ["MARS", "the rover"], size=20, bold=True, color=NAVY,
           align=CENTER)

    arrow_l = earth_l + earth_d + I(0.25)
    arrow_r = mars_l - I(0.25)
    _arrow(slide, arrow_l, mid_y, arrow_r, mid_y, color=TEAL, width=4,
           dashed=True)
    _label(slide, arrow_l, mid_y - I(0.95), arrow_r - arrow_l,
           ["A radio message takes", "3 to 22 MINUTES to get there"],
           size=22, bold=True, color=TEAL, align=CENTER)
    _label(slide, arrow_l, mid_y + I(0.2), arrow_r - arrow_l,
           "...and the answer takes just as long to come back",
           size=20, color=INK, align=CENTER)

    deck._note(slide,
               "So nobody can steer a Mars rover live. The team sends a PLAN, "
               "and the rover follows it all by itself.", "info")
    return slide


def mission_map(deck, title="Plan the route on the grid first", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "This is an EXAMPLE course. Yours on the floor may be laid out "
        "differently - tape a grid of one-foot squares, or use the floor "
        "tiles if your room has them.",
        "The dashed path is one way to do it. Ask for another. (Going "
        "round the bottom of the crater is just as good.)",
        "The big math idea: if the vehicle takes, say, 500 milliseconds to "
        "drive one square, how long for four squares? (2000.) Teams that "
        "measure one square first finish much faster than teams that guess.",
    ])

    cols, rows = 9, 5
    cell = I(0.74)
    grid_l = MARGIN_L + I(0.1)
    grid_t = BODY_TOP + I(0.5)
    for r in range(rows):
        for c in range(cols):
            _shape(slide, MSO_SHAPE.RECTANGLE, grid_l + cell * c,
                   grid_t + cell * r, cell, cell, fill=WHITE, edge=RULE,
                   edge_w=0.75)

    _box(slide, grid_l, grid_t + cell * 4, cell, cell, "GO",
         fill=START_GREEN, edge=NAVY, size=18, bold=True, color=NAVY,
         shape=MSO_SHAPE.RECTANGLE, edge_w=2)
    _box(slide, grid_l + cell * 8, grid_t, cell, cell, "", fill=SAMPLE,
         edge=NAVY, shape=MSO_SHAPE.RECTANGLE, edge_w=2, size=18)
    crater = _shape(slide, MSO_SHAPE.OVAL, grid_l + cell * 3 + I(0.1),
                    grid_t + cell * 1 + I(0.1), cell * 3 - I(0.2),
                    cell * 3 - I(0.2), fill=LIGHT_GREY, edge=GREY, edge_w=2)
    crater.name = "CRATER"
    _label(slide, grid_l + cell * 3, grid_t + cell * 2 + I(0.15), cell * 3,
           "CRATER", size=20, bold=True, color=GREY, align=CENTER)
    _label(slide, grid_l + cell * 6, grid_t - I(0.42), cell * 3,
           "SAMPLE", size=18, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)

    half = Emu(int(cell / 2))
    x0 = grid_l + half
    _arrow(slide, x0, grid_t + cell * 4 + I(0.1), x0, grid_t + half,
           color=TEAL, width=3.5, dashed=True)
    _arrow(slide, x0, grid_t + half, grid_l + cell * 8 + I(0.12),
           grid_t + half, color=TEAL, width=3.5, dashed=True)

    right = grid_l + cell * cols + I(0.45)
    width = MARGIN_L + CONTENT_W - right
    _label(slide, right, BODY_TOP + I(0.3), width, "One way to do it",
           size=24, bold=True, color=TEAL)
    _label(slide, right, BODY_TOP + I(0.9), width,
           ["forward 4 squares",
            "turn right",
            "forward 8 squares",
            "collect the sample!",
            "",
            "1 square = 1 foot",
            "How many milliseconds is ONE square? Measure it first."],
           size=20, color=INK)
    return slide


# ===================================================================
# Lesson 8
# ===================================================================

def switch_pad_map(deck, title="Your controller, and the name of every button",
                   speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Hold up a real controller next to the screen and match the two, "
        "button by button.",
        "The letter printed on the button is the letter in the code. Press "
        "A, and buttonA() is true. That's the whole rule.",
        "B is the brake button in the autopilot programs later on - B for "
        "Brake. Worth saying now so it sticks.",
        "If a controller from a different company has its letters in "
        "different places, trust the e00 test your teacher ran: it printed "
        "the name of each button as it was pressed.",
    ])

    body_l, body_t = I(3.55), BODY_TOP + I(1.2)
    body_w, body_h = I(6.2), I(2.9)
    pad = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, body_l, body_t, body_w,
                 body_h, fill=LIGHT_GREY, edge=NAVY, edge_w=2)
    pad.adjustments[0] = 0.35

    # Shoulders
    for x, name in ((body_l + I(0.55), "ZL"), (body_l + body_w - I(1.85), "ZR")):
        _box(slide, x, body_t - I(0.95), I(1.3), I(0.42), name, fill=WHITE,
             edge=NAVY, size=18, bold=True, color=NAVY)
    for x, name in ((body_l + I(0.55), "L"), (body_l + body_w - I(1.85), "R")):
        _box(slide, x, body_t - I(0.48), I(1.3), I(0.42), name, fill=WHITE,
             edge=NAVY, size=18, bold=True, color=NAVY)

    # Left stick, D-pad, right stick
    stick_d = I(0.9)
    left_stick = (body_l + I(0.75), body_t + I(0.35))
    _shape(slide, MSO_SHAPE.OVAL, left_stick[0], left_stick[1], stick_d,
           stick_d, fill=WHITE, edge=NAVY, edge_w=2)
    dpad = _shape(slide, MSO_SHAPE.CROSS, body_l + I(1.75), body_t + I(1.6),
                  I(0.9), I(0.9), fill=WHITE, edge=NAVY, edge_w=1.5)
    dpad.adjustments[0] = 0.33
    right_stick = (body_l + body_w - I(2.55), body_t + I(1.6))
    _shape(slide, MSO_SHAPE.OVAL, right_stick[0], right_stick[1], stick_d,
           stick_d, fill=WHITE, edge=NAVY, edge_w=2)

    # Face buttons: X top, Y left, A right, B bottom
    cx, cy = body_l + body_w - I(1.15), body_t + I(0.95)
    face_d = I(0.5)
    spots = {"X": (0, -1), "Y": (-1, 0), "A": (1, 0), "B": (0, 1)}
    for name, (dx, dy) in spots.items():
        x = cx + I(0.5 * dx) - Emu(int(face_d / 2))
        y = cy + I(0.5 * dy) - Emu(int(face_d / 2))
        _box(slide, x, y, face_d, face_d, name, fill=WHITE, edge=NAVY,
             size=18, bold=True, color=NAVY, shape=MSO_SHAPE.OVAL)

    # Callouts, left column
    left_w = I(2.85)
    _label(slide, MARGIN_L, BODY_TOP + I(0.05), left_w,
           ["L and ZL", "buttonL()  buttonZL()"], size=18, color=INK)
    _label(slide, MARGIN_L, body_t + I(0.3), left_w,
           ["LEFT STICK: drive", "leftStickX()", "leftStickY()"], size=18,
           color=INK)
    _label(slide, MARGIN_L, body_t + I(1.75), left_w,
           ["D-PAD", "dpadUp()  dpadDown()", "dpadLeft()  dpadRight()"],
           size=18, color=INK)
    _plain_line(slide, MARGIN_L + left_w - I(0.2), body_t + I(0.55),
                left_stick[0], left_stick[1] + Emu(int(stick_d / 2)),
                color=TEAL, width=1.5)
    _plain_line(slide, MARGIN_L + left_w - I(0.2), body_t + I(2.05),
                body_l + I(1.75), body_t + I(2.05), color=TEAL, width=1.5)

    # Callouts, right column
    right_l = body_l + body_w + I(0.35)
    right_w = MARGIN_L + CONTENT_W - right_l
    _label(slide, right_l, BODY_TOP + I(0.05), right_w,
           ["R and ZR", "buttonR()  buttonZR()"], size=18, color=INK)
    _label(slide, right_l, body_t + I(0.2), right_w,
           ["X  Y  A  B", "buttonA()  tappedA()", "...and the same for B, X, Y"],
           size=18, color=INK)
    _label(slide, right_l, body_t + I(1.75), right_w,
           ["RIGHT STICK", "rightStickX()", "rightStickY()"], size=18,
           color=INK)
    _plain_line(slide, right_l - I(0.05), body_t + I(1.05), cx + I(0.8),
                body_t + I(0.95), color=TEAL, width=1.5)
    _plain_line(slide, right_l - I(0.05), body_t + I(2.05),
                right_stick[0] + stick_d, right_stick[1] + Emu(int(stick_d / 2)),
                color=TEAL, width=1.5)

    deck._note(slide,
               "The letter PRINTED on the button is the letter in the code. "
               "Press A, and buttonA() is true.", "info")
    return slide


def if_else_flow(deck, title="if and else: the rover makes a choice",
                 speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Trace one trip through with your finger: start at the diamond, "
        "answer the question, follow YES or NO, do the job.",
        "Then say the important part: as soon as it finishes, loop() starts "
        "again and asks the question AGAIN - thousands of times every "
        "second. That's why the lights change the instant you let go.",
        "Match the drawing to the code on the right, line by line. The "
        "diamond is the if, the YES box is the first pair of braces, and "
        "the NO box is the else.",
    ])

    dia_w, dia_h = I(3.6), I(1.6)
    dia_l = MARGIN_L + I(1.7)
    dia_t = BODY_TOP + I(0.15)
    _box(slide, dia_l, dia_t, dia_w, dia_h, ["Is A held", "down?"],
         fill=LIGHT_AMBER, edge=NAVY, size=22, bold=True, color=NAVY,
         shape=MSO_SHAPE.DIAMOND)

    box_w, box_h = I(3.0), I(1.0)
    box_t = dia_t + dia_h + I(1.0)
    yes_l = MARGIN_L
    no_l = MARGIN_L + I(3.75)
    _box(slide, yes_l, box_t, box_w, box_h, ["front lights", "GREEN"],
         fill=LIGHT_TEAL, edge=NAVY, size=22, bold=True, color=NAVY)
    _box(slide, no_l, box_t, box_w, box_h, ["front lights", "OFF"],
         fill=WHITE, edge=NAVY, size=22, bold=True, color=NAVY)

    dia_mid_y = dia_t + Emu(int(dia_h / 2))
    _arrow(slide, dia_l, dia_mid_y, yes_l + Emu(int(box_w / 2)), box_t,
           color=TEAL, width=3)
    _arrow(slide, dia_l + dia_w, dia_mid_y, no_l + Emu(int(box_w / 2)), box_t,
           color=AMBER, width=3)
    _label(slide, yes_l, dia_mid_y + I(0.2), I(1.2), "YES", size=22,
           bold=True, color=TEAL)
    _label(slide, no_l + box_w - I(1.0), dia_mid_y + I(0.2), I(1.0), "NO",
           size=22, bold=True, color=AMBER, align=PP_ALIGN.RIGHT)
    _label(slide, MARGIN_L, box_t + box_h + I(0.2), I(6.75),
           "...then loop() starts again, and asks again.", size=20,
           color=GREY, align=CENTER)

    code_l = MARGIN_L + I(7.3)
    code_w = CONTENT_W - I(7.3)
    _box(slide, code_l, dia_t + I(0.2), code_w, I(2.3),
         ["if (buttonA()) {", "  setFrontLights(GREEN);", "} else {",
          "  setFrontLights(OFF);", "}"],
         fill=CODE_BG, edge=RULE, size=20, font=CODE_FONT,
         shape=MSO_SHAPE.RECTANGLE, edge_w=0.75, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP)
    _label(slide, code_l, dia_t + I(2.7), code_w,
           ["if means: when this is true...", "else means: otherwise..."],
           size=20, color=INK)
    return slide


# ===================================================================
# Lesson 9
# ===================================================================

def stick_numbers(deck, title="What the stick really sends", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "The stick isn't an on/off button. It's two numbers at once: how "
        "far up or down, and how far left or right.",
        "Minus just means the other direction - down instead of up, left "
        "instead of right. Younger students don't need negative numbers "
        "from math class to get this; a thermometer below zero is enough.",
        "The wobble zone is the small circle in the middle. A stick you let "
        "go of springs back NEAR zero, not exactly to zero. They'll see "
        "that for themselves on the Serial Monitor in a minute.",
    ])

    diameter = I(3.6)
    cx = MARGIN_L + I(2.9)
    cy = BODY_TOP + I(2.65)
    _shape(slide, MSO_SHAPE.OVAL, cx - Emu(int(diameter / 2)),
           cy - Emu(int(diameter / 2)), diameter, diameter, fill=WHITE,
           edge=NAVY, edge_w=2)
    reach = Emu(int(diameter / 2))
    _plain_line(slide, cx - reach, cy, cx + reach, cy, color=RULE, width=1.5)
    _plain_line(slide, cx, cy - reach, cx, cy + reach, color=RULE, width=1.5)
    wobble = I(0.75)
    _shape(slide, MSO_SHAPE.OVAL, cx - Emu(int(wobble / 2)),
           cy - Emu(int(wobble / 2)), wobble, wobble, fill=LIGHT_AMBER,
           edge=AMBER, edge_w=1.5)

    _label(slide, cx - I(1.2), cy - reach - I(0.75), I(2.4), "UP  +100",
           size=20, bold=True, color=TEAL, align=CENTER)
    _label(slide, cx - I(1.2), cy + reach + I(0.08), I(2.4), "DOWN  -100",
           size=20, bold=True, color=AMBER, align=CENTER)
    _label(slide, cx - reach - I(1.3), cy - I(0.65), I(1.2),
           ["LEFT", "-100"], size=20, bold=True, color=AMBER, align=CENTER)
    _label(slide, cx + reach + I(0.1), cy - I(0.65), I(1.3),
           ["RIGHT", "+100"], size=20, bold=True, color=TEAL, align=CENTER)

    right = MARGIN_L + I(6.6)
    width = CONTENT_W - I(6.6)
    _label(slide, right, BODY_TOP + I(0.2), width,
           ["leftStickY()", "up is plus, down is minus"], size=22,
           color=INK)
    _label(slide, right, BODY_TOP + I(1.2), width,
           ["leftStickX()", "right is plus, left is minus"], size=22,
           color=INK)
    _label(slide, right, BODY_TOP + I(2.3), width,
           ["Let go, and it SHOULD say 0...",
            "...but a real stick often says 2, or -3!"], size=22,
           bold=True, color=AMBER)
    _label(slide, right, BODY_TOP + I(3.4), width,
           ["The WOBBLE ZONE (orange circle):",
            "small numbers count as zero."], size=22, color=INK)
    return slide


# ===================================================================
# Lesson 10
# ===================================================================

def smart_light_states(deck, title="What the lights say", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Each little vehicle is seen from above, front at the top. The "
        "strip across the top is the front lights, the strip across the "
        "bottom is the back lights.",
        "Ask who has noticed a car's brake lights get BRIGHTER when it "
        "stops, or white lights come on at the back when it reverses. "
        "These are the same rules, and everybody on the road knows them.",
        "Turning right: orange on the RIGHT side, front and back. Point "
        "out that on our vehicle the right side is two separate pieces of "
        "the light loop - it's why setRightLights() exists.",
    ])

    W, D, R, O = LED_WHITE, LED_DIM, LED_RED, LED_ORANGE
    cases = [
        ("Driving forward", [W] * 8, [D] * 8, ["back: dim red", "tail lights"]),
        ("Stopped", [W] * 8, [R] * 8, ["back: bright red", "brake lights"]),
        ("Reversing", [W] * 8, [W] * 8, ["back: white", "backup lights"]),
        ("Turning right", [W] * 4 + [O] * 4, [D] * 4 + [O] * 4,
         ["right side: orange", "turn signal"]),
    ]
    gap = I(0.21)
    panel_w = Emu(int((CONTENT_W - gap * 3) / 4))
    top = BODY_TOP + I(0.1)
    for index, (name, front, back, words) in enumerate(cases):
        left = MARGIN_L + (panel_w + gap) * index
        cx = left + Emu(int(panel_w / 2))
        _box(slide, left, top, panel_w, I(0.55), name, fill=LIGHT_TEAL,
             edge=NAVY, size=22, bold=True, color=NAVY)
        _vehicle(slide, cx, top + I(1.15), I(1.75), I(1.75), front=front,
                 back=back)
        _label(slide, left, top + I(3.3), panel_w, words, size=20,
               color=INK, align=CENTER)

    deck._note(slide,
               "Headlights are white at the front in every picture. Only the "
               "BACK lights and the turn signals change.", "info")
    return slide


def nap_vs_clock(deck, title="Why not just use delay() to blink?",
                 speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Top row: each grey block is a delay(). While the rover is inside "
        "it, it's asleep. It doesn't read the stick, so if you steer during "
        "a nap, nothing happens until it wakes up.",
        "Bottom row: blinkIsOn() just glances at the clock and answers yes "
        "or no straight away, so loop() goes round thousands of times a "
        "second and catches your stick instantly.",
        "Act it out: you're the rover. A student 'steers' by pointing left "
        "or right. Close your eyes and count to three before every look. "
        "Then do it with your eyes open, glancing at a clock. The class "
        "will feel the difference.",
    ])

    track_l = MARGIN_L + I(0.1)
    track_w = CONTENT_W - I(0.2)
    row_h = I(0.65)

    row1_t = BODY_TOP + I(0.85)
    _label(slide, track_l, BODY_TOP + I(0.2), track_w,
           "delay(): the rover takes a NAP. It can't feel the stick while it "
           "sleeps.", size=20, bold=True, color=AMBER)
    x = track_l
    blocks = [(0.35, True), (2.7, False)] * 4
    scale = track_w / I(sum(w for w, _ in blocks))
    for width, awake in blocks:
        w = Emu(int(I(width) * scale))
        if awake:
            _shape(slide, MSO_SHAPE.RECTANGLE, x, row1_t, w, row_h, fill=TEAL,
                   edge=NAVY, edge_w=0.75)
        else:
            _box(slide, x, row1_t, w, row_h, "Zzz...  delay(400)",
                 fill=LIGHT_GREY, edge=RULE, size=18, color=GREY,
                 shape=MSO_SHAPE.RECTANGLE, edge_w=0.75)
        x += w

    row2_t = BODY_TOP + I(2.85)
    _label(slide, track_l, BODY_TOP + I(2.2), track_w,
           "blinkIsOn(): the rover GLANCES at the clock, and keeps checking.",
           size=20, bold=True, color=TEAL)
    count = 30
    w = Emu(int(track_w / count))
    for n in range(count):
        _shape(slide, MSO_SHAPE.RECTANGLE, track_l + w * n, row2_t, w, row_h,
               fill=TEAL if n % 2 == 0 else LIGHT_TEAL, edge=NAVY,
               edge_w=0.5)

    push_x = track_l + Emu(int(track_w * 0.18))
    _plain_line(slide, push_x, row1_t - I(0.05), push_x,
                row1_t + row_h + I(0.05), color=NAVY, width=2)
    _plain_line(slide, push_x, row2_t - I(0.05), push_x,
                row2_t + row_h + I(0.4), color=NAVY, width=2)
    _label(slide, push_x - I(1.6), row2_t + row_h + I(0.45), I(3.2),
           "you push the stick", size=20, bold=True, color=NAVY,
           align=CENTER)

    late_x = track_l + Emu(int(I(3.05) * scale))
    _arrow(slide, late_x + I(0.6), row1_t + row_h + I(0.55),
           late_x + I(0.12), row1_t + row_h + I(0.05), color=AMBER, width=2.5)
    _label(slide, late_x + I(0.65), row1_t + row_h + I(0.3), I(3.2),
           "notices here - too late!", size=18, bold=True, color=AMBER)

    deck._note(slide,
               "Grey = asleep. Teal = awake and reading the stick. Driving "
               "programs need to stay awake.", "info")
    return slide


# ===================================================================
# Lesson 11
# ===================================================================

def design_cycle(deck, title="How engineers build things", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "This is the engineering design process. Different companies draw "
        "it with five, six or seven steps; the steps matter less than the "
        "shape. It's a CIRCLE.",
        "Ask where most people want to jump straight to. (CREATE - typing "
        "the code.) Ask what goes wrong when you skip PLAN. (You build the "
        "wrong thing, fast.)",
        "And the most important arrow is the one from IMPROVE back round "
        "to ASK. Nobody gets it right first time - not NASA, not you, not "
        "me.",
    ])

    steps = [
        ("1  ASK", "What should it do?"),
        ("2  IMAGINE", "Lots of ideas"),
        ("3  PLAN", "Draw it, write it"),
        ("4  CREATE", "Write the code"),
        ("5  TEST", "Wheels up first!"),
        ("6  IMPROVE", "Fix ONE thing"),
    ]
    cx = MARGIN_L + Emu(int(CONTENT_W / 2))
    cy = BODY_TOP + I(2.45)
    rx, ry = I(3.75), I(1.85)
    box_w, box_h = I(2.95), I(0.95)
    centers = []
    for index in range(len(steps)):
        angle = math.radians(-90 + 60 * index)
        centers.append((cx + Emu(int(rx * math.cos(angle))),
                        cy + Emu(int(ry * math.sin(angle)))))
    for index, ((name, what), (x, y)) in enumerate(zip(steps, centers)):
        fill = LIGHT_AMBER if index == 3 else LIGHT_TEAL
        _box(slide, x - Emu(int(box_w / 2)), y - Emu(int(box_h / 2)), box_w,
             box_h, [name, what], fill=fill, edge=NAVY, size=20, color=NAVY)
    for index in range(len(steps)):
        (x1, y1), (x2, y2) = centers[index], centers[(index + 1) % len(steps)]
        # Shorten each arrow so it runs edge to edge, not center to center.
        fx, fy = x1 + (x2 - x1) * 0.36, y1 + (y2 - y1) * 0.36
        tx, ty = x1 + (x2 - x1) * 0.64, y1 + (y2 - y1) * 0.64
        _arrow(slide, Emu(int(fx)), Emu(int(fy)), Emu(int(tx)), Emu(int(ty)),
               color=TEAL, width=3)
    _label(slide, cx - I(1.6), cy - I(0.3), I(3.2),
           ["Go round as many", "times as you need"], size=20, bold=True,
           color=AMBER, align=CENTER)
    return slide


# ===================================================================
# Lesson 12
# ===================================================================

def mission_day_course(deck, title="The Mission Day course", speaker=None):
    slide = deck.blank(title, speaker=speaker or [
        "Set this up before families arrive, with tape and cups. Mark the "
        "no-signal zone with a dashed line of tape or a row of paper "
        "plates, so everybody can see where autopilot has to take over.",
        "Run the stations in order, one team at a time, with the rest of "
        "the class watching from behind the arena line. Watching other "
        "teams is half of the learning on this day.",
        "Adjust the course to your room. The four stations matter; the "
        "exact layout doesn't.",
    ])

    arena_l, arena_t = MARGIN_L + I(0.1), BODY_TOP + I(0.2)
    arena_w, arena_h = I(8.4), I(4.25)
    _shape(slide, MSO_SHAPE.RECTANGLE, arena_l, arena_t, arena_w, arena_h,
           fill=WHITE, edge=AMBER, edge_w=3)

    _box(slide, arena_l + I(0.2), arena_t + arena_h - I(1.05), I(1.5), I(0.85),
         ["START", "FINISH"], fill=START_GREEN, edge=NAVY, size=18,
         bold=True, color=NAVY)

    # 1 Slalom: cups along the bottom
    for n in range(4):
        _shape(slide, MSO_SHAPE.OVAL, arena_l + I(2.4) + I(1.1) * n,
               arena_t + arena_h - I(0.75), I(0.36), I(0.36),
               fill=LIGHT_AMBER, edge=AMBER, edge_w=1.5)
    # 2 No-signal zone: dashed box up the right side
    _shape(slide, MSO_SHAPE.RECTANGLE, arena_l + arena_w - I(2.2),
           arena_t + I(0.25), I(1.95), I(2.5), fill=LIGHT_GREY, edge=NAVY,
           edge_w=2, dashed=True)
    _label(slide, arena_l + arena_w - I(2.2), arena_t + I(0.95), I(1.95),
           ["NO", "SIGNAL"], size=18, bold=True, color=NAVY, align=CENTER)
    # 3 Signal check: a corner, top middle
    _shape(slide, MSO_SHAPE.RIGHT_ARROW, arena_l + I(3.3), arena_t + I(0.45),
           I(1.4), I(0.6), fill=LED_ORANGE, edge=NAVY, edge_w=1)
    # 4 Parking: garage, top left
    _box(slide, arena_l + I(0.3), arena_t + I(0.3), I(1.3), I(1.1),
         "GARAGE", fill=WHITE, edge=TEAL, size=18, bold=True, color=TEAL,
         edge_w=2.5)

    badges = [
        (arena_l + I(4.1), arena_t + arena_h - I(1.45), "1"),
        (arena_l + arena_w - I(1.5), arena_t + I(2.85), "2"),
        (arena_l + I(3.75), arena_t + I(1.15), "3"),
        (arena_l + I(1.75), arena_t + I(0.55), "4"),
    ]
    for x, y, number in badges:
        _box(slide, x, y, I(0.55), I(0.55), number, fill=NAVY, edge=NAVY,
             size=20, bold=True, color=WHITE, shape=MSO_SHAPE.OVAL)

    right = MARGIN_L + I(8.85)
    width = CONTENT_W - I(8.85)
    _label(slide, right, BODY_TOP + I(0.2), width,
           ["1  Slalom", "weave through the cups", "",
            "2  No-signal zone", "press X: autopilot across", "",
            "3  Signal check", "turn signal on as you turn", "",
            "4  Precision parking", "all four wheels in the garage"],
           size=18, color=INK)
    return slide
